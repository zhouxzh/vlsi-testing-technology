"""Convert slide text and speaker notes from the local chapter decks to Markdown."""

from __future__ import annotations

import argparse
import json
import re
import statistics
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


CHAPTER_RE = re.compile(r"^第(?P<number>\d+)章\.pptx$")


def natural_chapter(path: Path) -> tuple[int, str]:
    match = CHAPTER_RE.match(path.name)
    if not match:
        raise ValueError(f"Unexpected chapter filename: {path.name}")
    return int(match.group("number")), path.stem


def shape_top(shape) -> tuple[float, float]:
    return (float(shape.top or 0), float(shape.left or 0))


def unique_nonempty(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        value = value.strip()
        if not value or value in seen:
            continue
        seen.add(value)
        result.append(value)
    return result


def shape_texts(shape) -> list[str]:
    texts: list[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in sorted(shape.shapes, key=shape_top):
            texts.extend(shape_texts(child))
        return texts

    if getattr(shape, "has_text_frame", False):
        paragraph_texts = []
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs).strip()
            if text:
                paragraph_texts.append(text)
        if paragraph_texts:
            texts.append("\n".join(paragraph_texts))

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = unique_nonempty(cell.text for cell in row.cells)
            if cells:
                texts.append(" | ".join(cells))
    return texts


def slide_title(slide, extracted: list[str]) -> str:
    title_shape = getattr(slide.shapes, "title", None)
    if title_shape is not None and getattr(title_shape, "has_text_frame", False):
        title = title_shape.text.strip()
        if title:
            return " ".join(title.split())
    if extracted:
        return " ".join(extracted[0].split())[:80]
    return "（无标题）"


def markdown_fence(text: str) -> str:
    longest = max((len(match.group(0)) for match in re.finditer(r"`+", text)), default=0)
    fence = "`" * max(3, longest + 1)
    return f"{fence}text\n{text}\n{fence}"


def write_text_lf(path: Path, text: str) -> None:
    path.write_bytes(text.encode("utf-8"))


def extract_deck(path: Path, output_root: Path, chapter: int) -> dict:
    deck = Presentation(path)
    slug = f"ch{chapter:02d}"
    asset_rel = Path("assets") / slug
    slide_lines: list[str] = []
    slide_titles: list[str] = []
    page_anchors: list[str] = []
    text_lengths: list[int] = []
    page_text_chars = 0
    page_note_chars = 0
    pages_with_notes = 0

    for index, slide in enumerate(deck.slides, start=1):
        texts = unique_nonempty(
            text
            for shape in sorted(slide.shapes, key=shape_top)
            for text in shape_texts(shape)
        )
        title = slide_title(slide, texts)
        slide_titles.append(title)
        page_anchors.append(f"{slug}-s{index:02d}")
        image_name = f"page-{index:03d}.jpg"
        image_path = asset_rel / image_name

        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            notes = re.sub(r"\n{3,}", "\n\n", notes)
            if notes:
                pages_with_notes += 1
                page_note_chars += len(notes)

        visible_text = "\n\n".join(texts) if texts else "本页以图像、公式或结构图为主。"
        text_length = sum(len(text) for text in texts)
        text_lengths.append(text_length)
        page_text_chars += text_length
        slide_lines.extend(
            [
                f"## 第 {index} 页 · {title} {{#{page_anchors[-1]}}}",
                "",
                f"![第 {index} 页](<{image_path.as_posix()}>)",
                "",
                "### 页面文字",
                "",
                markdown_fence(visible_text),
                "",
            ]
        )
        if notes:
            slide_lines.extend(["### 讲师备注", "", markdown_fence(notes), ""])

    source_name = path.name
    title = slide_titles[0] if slide_titles else path.stem
    document = "\n".join(
        [
            "---",
            f'title: "{path.stem}素材"',
            f"source: \"{source_name}\"",
            "source_type: PowerPoint",
            f"pages: {len(deck.slides)}",
            "---",
            "",
            f"# {path.stem}素材",
            "",
            "> 本页由本地原始课件自动整理。原始 PPT 与 PDF 不进入版本库；",
            "> 页面图像、文字和讲师备注仅作为本课程备课与授课素材。",
            "",
            "## 本章素材导航",
            "",
            *[
                f"- [第 {i} 页 · {name}](#{anchor})"
                for i, (name, anchor) in enumerate(
                    zip(slide_titles, page_anchors, strict=True), start=1
                )
            ],
            "",
            *slide_lines,
            "## 使用边界",
            "",
            "- 图像用于保留公式、波形、结构图和版式信息，文字抽取不能替代原页。",
            "- 讲授前应核对数值、单位、器件条件和引用来源。",
            "- 本文件不替代原始教材、论文或器件手册。",
            "",
        ]
    )
    write_text_lf(output_root / f"{slug}.md", document)

    return {
        "chapter": chapter,
        "title": title,
        "source": source_name,
        "output": f"{slug}.md",
        "assets": asset_rel.as_posix(),
        "pages": len(deck.slides),
        "pages_with_notes": pages_with_notes,
        "text_characters": page_text_chars,
        "note_characters": page_note_chars,
        "median_text_characters": statistics.median(text_lengths),
    }


def write_index(output_root: Path, entries: list[dict]) -> None:
    total_pages = sum(item["pages"] for item in entries)
    total_notes = sum(item["pages_with_notes"] for item in entries)
    lines = [
        "---",
        "title: 十二份原始课件素材",
        "---",
        "",
        "# 十二份原始课件素材",
        "",
        "本目录把 `reference/` 中的 12 份本地 PPT 转换为可检索的 Markdown。",
        "`reference/` 已被 Git 忽略，原文和 PDF 不会随课程主页发布。",
        "",
        "| 素材 | 识别标题 | 页数 | 含讲师备注页 |",
        "| --- | --- | ---: | ---: |",
    ]
    for item in entries:
        title = item["title"].replace("|", "\\|")
        lines.append(
            f"| [{item['chapter']:02d}]({item['output']}) | {title} | "
            f"{item['pages']} | {item['pages_with_notes']} |"
        )
    lines.extend(
        [
            "",
            f"共 **{len(entries)} 份** 课件、**{total_pages} 页**，"
            f"其中 **{total_notes} 页**带有可用讲师备注。",
            "",
            "## 使用方式",
            "",
            "1. 先按页面图像确认图表、公式和版式，再用文字区检索关键词。",
            "2. 课堂课件只选取支撑 8 章教学目标的内容，不逐页搬运原始 PPT。",
            "3. 进入正式讲稿前核对器件条件、单位、缩写和图表来源。",
            "",
        ]
    )
    write_text_lf(output_root / "README.md", "\n".join(lines))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--reference", type=Path, default=Path("reference"))
    parser.add_argument("--output", type=Path, default=Path("reference/materials-v1"))
    args = parser.parse_args()

    args.output.mkdir(parents=True, exist_ok=True)
    deck_paths = sorted(
        (path for path in args.reference.glob("第*章.pptx") if CHAPTER_RE.match(path.name)),
        key=lambda path: natural_chapter(path)[0],
    )
    if len(deck_paths) != 12:
        raise SystemExit(f"Expected 12 chapter decks, found {len(deck_paths)}")

    entries = []
    for path in deck_paths:
        chapter, _ = natural_chapter(path)
        entry = extract_deck(path, args.output, chapter)
        entries.append(entry)
        print(
            f"{path.name}: {entry['pages']} pages, "
            f"{entry['text_characters']} text chars, "
            f"{entry['pages_with_notes']} notes pages"
        )

    write_index(args.output, entries)
    write_text_lf(
        args.output / "manifest.json",
        json.dumps(entries, ensure_ascii=False, indent=2) + "\n",
    )
    print(f"Wrote {len(entries)} Markdown files to {args.output}")


if __name__ == "__main__":
    main()
