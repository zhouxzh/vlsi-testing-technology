"""Extract the instructor's eFlash and IGBT source material into local Markdown."""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from docx import Document
from PIL import Image
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / "reference"
OUTPUT = REFERENCE / "chapter1-cases"
EFLASH_PPTX = REFERENCE / "55nm工艺节点eflash存算一体AI芯片.pptx"
EFLASH_SLIDES = OUTPUT / "eflash-slides"
EFLASH_FLOW = REFERENCE / "流片图.png"
IGBT_DIR = REFERENCE / "2024-2-IGBT"


def write_text_lf(path: Path, text: str) -> None:
    path.write_bytes(text.encode("utf-8"))


def fence(text: str) -> str:
    longest = max((len(match.group(0)) for match in re.finditer(r"`+", text)), default=0)
    marker = "`" * max(3, longest + 1)
    return f"{marker}text\n{text}\n{marker}"


def slide_texts(slide) -> list[str]:
    texts: list[str] = []
    for shape in sorted(slide.shapes, key=lambda item: (item.top or 0, item.left or 0)):
        if getattr(shape, "has_text_frame", False):
            value = shape.text.strip()
            if value and value not in texts:
                texts.append(value)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    value = " | ".join(cells)
                    if value not in texts:
                        texts.append(value)
    return texts


def extract_pptx() -> dict:
    deck = Presentation(EFLASH_PPTX)
    rendered = sorted(EFLASH_SLIDES.glob("page-*.jpg"))
    if len(rendered) != len(deck.slides):
        raise SystemExit(
            "eFlash slide images are missing. Run "
            "scripts/export_single_pptx.ps1 first."
        )
    target_dir = OUTPUT / "eflash-ai-chip"
    target_dir.mkdir(parents=True, exist_ok=True)
    lines = [
        "---",
        "title: 55 nm eFlash 存算一体 AI 芯片案例素材",
        f'source: "{EFLASH_PPTX.name}"',
        f"pages: {len(deck.slides)}",
        "---",
        "",
        "# 55 nm eFlash 存算一体 AI 芯片案例素材",
        "",
        "> 本文件由授课教师本地 PPT 整理，仅供课程备课，不随网站发布。",
        "",
    ]
    for index, slide in enumerate(deck.slides, start=1):
        texts = slide_texts(slide)
        title = " ".join(texts[0].split())[:100] if texts else "（无文字标题）"
        lines.extend(
            [
                f"## 第 {index} 页 · {title} {{#slide-{index:03d}}}",
                "",
                f"![第 {index} 页](../eflash-slides/page-{index:03d}.jpg)",
                "",
                fence("\n\n".join(texts) if texts else "本页主要为图像。"),
                "",
            ]
        )
    write_text_lf(target_dir / "README.md", "\n".join(lines))
    if EFLASH_FLOW.is_file():
        shutil.copy2(EFLASH_FLOW, target_dir / "流片图.png")
    with Image.open(EFLASH_FLOW) as image:
        dimensions = list(image.size)
    return {
        "source": EFLASH_PPTX.name,
        "slides": len(deck.slides),
        "markdown": str((target_dir / "README.md").relative_to(OUTPUT)),
        "tapeout_image": "流片图.png" if EFLASH_FLOW.is_file() else None,
        "tapeout_image_size": dimensions,
    }


def extract_docx() -> dict:
    source = IGBT_DIR / "IGBT.docx"
    document = Document(source)
    target_dir = OUTPUT / "igbt-power-device"
    target_dir.mkdir(parents=True, exist_ok=True)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    tables: list[list[list[str]]] = []
    for table in document.tables:
        rows = [
            [cell.text.strip().replace("\n", " ") for cell in row.cells]
            for row in table.rows
        ]
        if rows:
            tables.append(rows)
    lines = [
        "---",
        "title: IGBT 与功率器件流片测试案例素材",
        f'source: "{source.name}"',
        "---",
        "",
        "# IGBT 与功率器件流片测试案例素材",
        "",
        "> 本文件由授课教师本地 Word 文档整理，仅供课程备课，不随网站发布。",
        "",
        "## 文档文字",
        "",
        fence("\n\n".join(paragraphs) if paragraphs else "文档以图片为主。"),
        "",
        "## 表格",
        "",
    ]
    if tables:
        for index, rows in enumerate(tables, start=1):
            width = max(len(row) for row in rows)
            normalized = [row + [""] * (width - len(row)) for row in rows]
            headers = [cell.replace("|", "\\|") for cell in normalized[0]]
            lines.extend(
                [
                    f"### 表 {index}",
                    "",
                    "| " + " | ".join(headers) + " |",
                    "| " + " | ".join("---" for _ in headers) + " |",
                    *[
                        "| "
                        + " | ".join(cell.replace("|", "\\|") for cell in row)
                        + " |"
                        for row in normalized[1:]
                    ],
                    "",
                ]
            )
    else:
        lines.extend(["文档中没有可提取的表格。", ""])
    lines.extend(["## 图片", ""])
    copied = []
    for name in ("IGBT.png", "driver.png"):
        source_image = IGBT_DIR / name
        if source_image.is_file():
            shutil.copy2(source_image, target_dir / name)
            copied.append(name)
            lines.extend([f"### {name}", "", f"![{name}]({name})", ""])
    write_text_lf(target_dir / "README.md", "\n".join(lines))
    return {
        "source": source.name,
        "paragraphs": len(paragraphs),
        "markdown": str((target_dir / "README.md").relative_to(OUTPUT)),
        "images": copied,
    }


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    eflash = extract_pptx()
    igbt = extract_docx()
    write_text_lf(
        OUTPUT / "manifest.json",
        json.dumps(
            {"eflash": eflash, "igbt": igbt},
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
    )
    print(json.dumps({"eflash": eflash, "igbt": igbt}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
